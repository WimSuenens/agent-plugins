"""Build a .pptx from a deck spec.

    build_deck.py work/spec.json [--out build/deck.pptx] [--no-embed]

The spec carries the measurements; this carries none. Coordinates may be given
in source pixels (`*_px`, resolved through the slide's `px_per_in` and an
optional band transform) or directly in inches (`*_in`). See
references/spec-schema.md for every element type and field.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import draw as D  # noqa: E402
import iconlib  # noqa: E402
import package  # noqa: E402

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT,
         "j": PP_ALIGN.JUSTIFY}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


# --- fonts -------------------------------------------------------------------

class Fonts:
    """Resolves family+weight to a file, and reads the metrics off that file.

    Baseline placement depends on the font's own ascent, so reading it rather
    than assuming one keeps text on its measured baseline whatever family the
    source turns out to use.
    """

    def __init__(self, spec_dir: Path, cfg: dict):
        self.dir = (spec_dir / cfg.get("dir", "fonts")).resolve()
        self.families: dict[str, dict[str, str]] = cfg.get("families", {})
        self.embed = cfg.get("embed", True)
        self._metrics: dict[str, tuple[float, float]] = {}

    def file(self, family: str, bold: bool = False, italic: bool = False) -> Path | None:
        styles = self.families.get(family)
        if not styles:
            return None
        for key in ([("boldItalic" if bold else "italic")] if italic else []) + \
                   (["bold"] if bold else []) + ["regular"]:
            if key in styles:
                path = self.dir / styles[key]
                if path.exists():
                    return path
        return None

    def metrics(self, family: str, bold: bool = False) -> tuple[float, float]:
        key = f"{family}|{bold}"
        if key not in self._metrics:
            self._metrics[key] = (1.0, 0.2)
            path = self.file(family, bold)
            if path:
                try:
                    from fontTools.ttLib import TTFont

                    f = TTFont(str(path))
                    upem = f["head"].unitsPerEm
                    self._metrics[key] = (f["hhea"].ascender / upem,
                                          -f["hhea"].descender / upem)
                except Exception:
                    pass
        return self._metrics[key]

    def embed_map(self) -> dict[str, dict[str, Path]]:
        out: dict[str, dict[str, Path]] = {}
        for family, styles in self.families.items():
            files = {s: self.dir / f for s, f in styles.items() if (self.dir / f).exists()}
            if files:
                out[family] = files
        return out


# --- coordinates -------------------------------------------------------------

class Geometry:
    """Resolves spec coordinates to inches."""

    def __init__(self, slide_spec: dict, defaults: dict):
        self.ppi = slide_spec.get("px_per_in") or defaults.get("px_per_in") or 96.0
        self.bands = slide_spec.get("bands", {})

    def _band(self, name: str | None):
        return self.bands.get(name) if name else None

    def x(self, px: float) -> float:
        return px / self.ppi

    def y(self, px: float, band: str | None) -> float:
        b = self._band(band)
        if not b:
            return px / self.ppi
        return b["origin_in"] + (px - b["src0"]) * b.get("scale", 1.0) / self.ppi

    def dy(self, px: float, band: str | None) -> float:
        b = self._band(band)
        return px * (b.get("scale", 1.0) if b else 1.0) / self.ppi

    def box(self, el: dict, band: str | None, art_aspect: float | None = None):
        """-> (x, y, w, h) in inches, from box_px [x0,y0,x1,y1] or box_in [x,y,w,h].

        With `keep_aspect` the width is recomputed from the height and the
        artwork's own aspect. A band compresses the y-axis but not the x-axis,
        so anything with fixed proportions -- a logo, a map, a pictogram --
        would otherwise come out subtly stretched.
        """
        if "box_in" in el:
            x, y, w, h = el["box_in"]
        else:
            x0, y0, x1, y1 = el["box_px"]
            x, y = self.x(x0), self.y(y0, band)
            w, h = self.x(x1 - x0), self.dy(y1 - y0, band)
        anchor = el.get("keep_aspect")
        if anchor and art_aspect:
            new_w = h * art_aspect
            if anchor == "center":
                x += (w - new_w) / 2
            elif anchor == "right":
                x += w - new_w
            w = new_w
        return (x, y, w, h)

    def val(self, el: dict, key: str, band: str | None = None, axis: str = "x",
            default=None):
        """Read `key_px` or `key_in` from an element."""
        if f"{key}_in" in el:
            return el[f"{key}_in"]
        if f"{key}_px" in el:
            px = el[f"{key}_px"]
            return self.x(px) if axis == "x" else (
                self.y(px, band) if axis == "y" else self.dy(px, band))
        return default


# --- element drawing ---------------------------------------------------------

def runs_of(el_runs, defaults, fonts: Fonts):
    out = []
    for r in el_runs:
        out.append({
            "text": r["t"],
            "size": r.get("pt", defaults.get("pt", 18)),
            "font": r.get("font", defaults.get("font", "Arial")),
            "bold": bool(r.get("b", False)),
            "italic": bool(r.get("i", False)),
            "color": r.get("c", defaults.get("color", "000000")),
        })
    return out


def radius_adj(el: dict, geo: Geometry, w: float, h: float):
    """roundRect adj is a fraction of the shorter side, not an absolute radius."""
    if "radius" in el:
        return el["radius"]
    r_in = geo.val(el, "radius", None, "len")
    if r_in is None:
        return None
    return max(0.0, min(0.5, r_in / max(min(w, h), 1e-6)))


def draw_element(shapes, el: dict, geo: Geometry, defaults: dict, fonts: Fonts,
                 traced: dict, spec_dir: Path):
    kind = el["type"]
    band = el.get("band")
    name = el.get("name")

    if kind in ("rect", "ellipse"):
        x, y, w, h = geo.box(el, band)
        if kind == "ellipse":
            shp = D.ellipse(shapes, x, y, w, h, fill=el.get("fill"),
                            line=el.get("line"), line_pt=el.get("line_pt", 1.0), name=name)
        else:
            shp = D.rect(shapes, x, y, w, h, radius=radius_adj(el, geo, w, h),
                         fill=el.get("fill"), gradient=el.get("gradient"),
                         line=el.get("line"), line_pt=el.get("line_pt", 1.0), name=name)
        if el.get("shadow"):
            D.soft_shadow(shp, **(el["shadow"] if isinstance(el["shadow"], dict) else {}))
        return shp

    if kind == "line":
        x1, y1 = el["from_px"] if "from_px" in el else el["from_in"]
        x2, y2 = el["to_px"] if "to_px" in el else el["to_in"]
        if "from_px" in el:
            x1, y1 = geo.x(x1), geo.y(y1, band)
            x2, y2 = geo.x(x2), geo.y(y2, band)
        shp = D.line_shape(shapes, x1, y1, x2, y2, color=el.get("color", "000000"),
                           width_pt=el.get("width_pt", 1.0), name=name)
        if el.get("arrow"):
            D.arrow_end(shp)
        return shp

    if kind == "picture":
        art = None
        if el.get("keep_aspect"):
            from PIL import Image

            with Image.open((spec_dir / el["file"]).resolve()) as im:
                art = im.width / im.height
        x, y, w, h = geo.box(el, band, art)
        pic = shapes.add_picture(str((spec_dir / el["file"]).resolve()),
                                 Inches(x), Inches(y), Inches(w), Inches(h))
        pic.name = name or Path(el["file"]).stem
        r = radius_adj(el, geo, w, h)
        if r:
            set_prst_geom(pic, "roundRect", r)
        if el.get("shadow"):
            D.soft_shadow(pic, **(el["shadow"] if isinstance(el["shadow"], dict) else {}))
        return pic

    if kind == "text":
        runs = runs_of(el["runs"], defaults, fonts)
        lead = max(runs, key=lambda r: r["size"])
        ascent, descent = fonts.metrics(lead["font"], lead["bold"])
        left = geo.val(el, "left", band, "x", 0.0)
        width = geo.val(el, "width", band, "len", 6.0)
        baseline = geo.val(el, "baseline", band, "y")
        return D.text_line(shapes, runs, left=left, width=width, baseline=baseline,
                           align=ALIGN.get(el.get("align", "c"), PP_ALIGN.CENTER),
                           name=name, ascent=ascent, descent=descent)

    if kind == "textbox":
        x, y, w, h = geo.box(el, band)
        paras = [runs_of(p, defaults, fonts) for p in el["paragraphs"]]
        return D.textbox(shapes, paras, left=x, top=y, width=w, height=h,
                         align=ALIGN.get(el.get("align", "l"), PP_ALIGN.LEFT),
                         anchor=ANCHOR.get(el.get("anchor", "t"), MSO_ANCHOR.TOP),
                         line_spacing=(Pt(el["line_spacing_pt"])
                                       if el.get("line_spacing_pt") else None),
                         space_after=el.get("space_after_pt", 0), name=name)

    if kind == "bullets":
        bullet = el.get("bullet", {})
        ascent, descent = fonts.metrics(el.get("font", defaults.get("font", "Arial")))
        return D.bullet_list(
            shapes, el["items"],
            left=geo.val(el, "left", band, "x", 0.0),
            width=geo.val(el, "width", band, "len", 5.0),
            first_baseline=geo.val(el, "first_baseline", band, "y"),
            spacing=geo.val(el, "spacing", band, "len", 0.3),
            size=el.get("pt", defaults.get("pt", 16)),
            font=el.get("font", defaults.get("font", "Arial")),
            color=el.get("color", defaults.get("color", "000000")),
            bullet_color=bullet.get("color", el.get("color", "000000")),
            text_offset=geo.val(bullet, "offset", band, "x", 0.3),
            char=bullet.get("char", "ü"), bullet_font=bullet.get("font", "Wingdings"),
            size_pct=bullet.get("size_pct", 120), ascent=ascent, descent=descent,
            name=name)

    if kind == "table":
        x, y, w, h = geo.box(el, band)
        return D.table(shapes, el["rows"], left=x, top=y, width=w, height=h,
                       font=el.get("font", defaults.get("font", "Arial")),
                       size=el.get("pt", defaults.get("pt", 12)),
                       color=el.get("color", defaults.get("color", "000000")),
                       header_fill=el.get("header_fill"),
                       header_color=el.get("header_color"),
                       body_fill=el.get("body_fill"),
                       col_widths=el.get("col_widths"), name=name)

    if kind == "icon":
        prims = el.get("prims") or iconlib.icon(el["icon"])
        return D.draw_icon(
            shapes, prims,
            center_x=geo.val(el, "center_x", band, "x"),
            top=geo.val(el, "top", band, "y"),
            height=geo.val(el, "height", band, "len"),
            color=el.get("color", "000000"),
            stroke_pt=el.get("stroke_pt", 1.9),
            name=name or f"Icon {el.get('icon', 'custom')}")

    if kind == "freeform":
        d, src = resolve_path(el, traced)
        art = (src[2] - src[0]) / max(src[3] - src[1], 1e-6)
        x, y, w, h = geo.box(el, band, art)
        return D.freeform(shapes, d, x, y, w, h, src, fill=el.get("fill"),
                          color=el.get("line"), width_pt=el.get("line_pt", 1.5),
                          name=name)

    if kind == "group":
        members = [draw_element(shapes, child, geo, defaults, fonts, traced, spec_dir)
                   for child in el["children"]]
        return D.group(shapes, members, name or "Group")

    raise ValueError(f"unknown element type: {kind!r}")


def resolve_path(el: dict, traced: dict):
    """Path data either inline (`d`) or by reference into traced.json (`d_ref`)."""
    if "d" in el:
        d = el["d"]
        src = tuple(el["src"]) if "src" in el else None
    else:
        key = el["d_ref"].split("#")[-1]
        entry = traced[key]
        d, src = entry["d"], tuple(entry["bbox"])
    if src is None:
        from custgeom import bbox, parse_path

        src = bbox(parse_path(d))
    return d, src


def set_prst_geom(shape, prst: str, adj: float) -> None:
    """Give a picture rounded corners without converting it into a shape."""
    from lxml import etree
    from pptx.oxml.ns import qn

    sp_pr = shape._element.spPr
    old = sp_pr.find(qn("a:prstGeom"))
    new = etree.fromstring(
        '<a:prstGeom xmlns:a="{ns}" prst="{p}"><a:avLst>'
        '<a:gd name="adj" fmla="val {v}"/></a:avLst></a:prstGeom>'.format(
            ns=D.A_NS, p=prst, v=int(adj * 100000)))
    if old is not None:
        sp_pr.replace(old, new)
    else:
        sp_pr.insert(0, new)


# --- build -------------------------------------------------------------------

def build(spec_path: Path, out: Path | None = None, embed: bool | None = None) -> Path:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    spec_dir = spec_path.parent
    deck = spec.get("deck", {})
    defaults = spec.get("defaults", {})
    fonts = Fonts(spec_dir, deck.get("fonts", {}))

    traced: dict = {}
    for ref in ([deck["traced"]] if isinstance(deck.get("traced"), str)
                else deck.get("traced", [])):
        traced.update(json.loads((spec_dir / ref).read_text()))

    prs = Presentation()
    prs.slide_width = Inches(deck.get("slide_w_in", 13.3333))
    prs.slide_height = Inches(deck.get("slide_h_in", 7.5))
    blank = prs.slide_layouts[6]

    counts = {"slides": 0, "elements": 0}
    for slide_spec in spec["slides"]:
        slide = prs.slides.add_slide(blank)
        if slide_spec.get("background"):
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = D.hexrgb(slide_spec["background"])
        geo = Geometry(slide_spec, defaults)
        for el in slide_spec["elements"]:
            draw_element(slide.shapes, el, geo, defaults, fonts, traced, spec_dir)
            counts["elements"] += 1
        counts["slides"] += 1

    out = out or (spec_dir / deck.get("output", "deck.pptx"))
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".raw.pptx")
    prs.save(tmp)

    theme = deck.get("theme", {})
    colours = {k: v for k, v in theme.items() if k.startswith("accent") or k in ("dk1", "lt1", "dk2", "lt2")}
    should_embed = fonts.embed if embed is None else embed
    notes = package.process(
        tmp, out, colours=colours,
        major=theme.get("major_font"), minor=theme.get("minor_font"),
        families=fonts.embed_map() if should_embed else None)

    for note in notes:
        print(" ", note)
    print(f"wrote {out}  ({out.stat().st_size:,} bytes, "
          f"{counts['slides']} slide(s), {counts['elements']} elements)")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("--out")
    ap.add_argument("--no-embed", action="store_true")
    args = ap.parse_args()
    build(Path(args.spec), Path(args.out) if args.out else None,
          embed=False if args.no_embed else None)
    return 0


if __name__ == "__main__":
    sys.exit(main())
