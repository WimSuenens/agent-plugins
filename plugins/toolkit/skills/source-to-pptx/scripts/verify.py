"""Inventory a built deck: what is text, what is a shape, what is a picture.

    verify.py build/deck.pptx [--expect-pictures 1] [--quiet]

This is the editability check. A rebuild can look perfect and still be a failure
-- if the text came out as pictures, or an icon as a bitmap, the deck cannot be
edited and the whole exercise was pointless. Run it before handing anything over,
and read the totals: text frames should roughly match the number of text runs on
the page, and pictures should match the number of actual photographs.
"""
from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def walk(shapes, depth=0, quiet=False):
    counts = {"text": 0, "shape": 0, "picture": 0, "group": 0, "connector": 0, "table": 0}
    for shp in shapes:
        kind = "shape"
        stype = str(shp.shape_type)
        if "PICTURE" in stype:
            kind = "picture"
        elif getattr(shp, "has_table", False):
            kind = "table"
        elif shp.has_text_frame and shp.text_frame.text.strip():
            kind = "text"
        if "GROUP" in stype:
            kind = "group"
        elif shp._element.tag.endswith("}cxnSp"):
            kind = "connector"
        counts[kind] += 1

        if not quiet:
            extra = ""
            if kind == "text":
                runs = [r for p in shp.text_frame.paragraphs for r in p.runs]
                fonts = sorted({r.font.name for r in runs if r.font.name})
                sizes = sorted({round(r.font.size.pt, 1) for r in runs if r.font.size})
                text = shp.text_frame.text.replace("\n", " / ")[:44]
                extra = f'  "{text}"  {fonts} {sizes}pt'
            elif kind in ("shape", "connector"):
                if shp._element.spPr.find(A_NS + "custGeom") is not None:
                    extra = "  custGeom"
            print(f"{'  ' * depth}- {shp.name:44s} [{kind}]{extra}")

        if kind == "group":
            for k, v in walk(shp.shapes, depth + 1, quiet).items():
                counts[k] += v
    return counts


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--expect-pictures", type=int,
                    help="fail if the deck has a different number of pictures")
    ap.add_argument("--min-text", type=int, default=1,
                    help="fail below this many text frames (default 1)")
    ap.add_argument("--quiet", action="store_true", help="totals only")
    args = ap.parse_args()

    deck = Path(args.deck)
    prs = Presentation(deck)
    print(f"{deck.name}: {Emu(prs.slide_width).inches:.3f} x "
          f"{Emu(prs.slide_height).inches:.3f} in, {len(prs.slides)} slide(s)\n")

    totals = {k: 0 for k in ("text", "shape", "picture", "group", "connector", "table")}
    for i, slide in enumerate(prs.slides, start=1):
        if not args.quiet:
            print(f"--- slide {i} ---")
        for k, v in walk(slide.shapes, quiet=args.quiet).items():
            totals[k] += v
    print("\ntotals: " + ", ".join(f"{k}={v}" for k, v in totals.items()))

    with zipfile.ZipFile(deck) as z:
        names = z.namelist()
        fonts = [n for n in names if n.startswith("ppt/fonts/")]
        media = [n for n in names if n.startswith("ppt/media/")]
        pres = z.read("ppt/presentation.xml").decode()
        theme = z.read("ppt/theme/theme1.xml").decode()
    families = re.findall(r'<p:font typeface="([^"]+)"', pres)
    accents = re.findall(r'<a:accent[12]>\s*<a:srgbClr val="([0-9A-Fa-f]{6})"', theme)
    major = re.search(r'<a:majorFont>\s*<a:latin typeface="([^"]+)"', theme)
    minor = re.search(r'<a:minorFont>\s*<a:latin typeface="([^"]+)"', theme)
    print(f"embedded font parts: {len(fonts)}  families: {families}")
    print(f"media parts: {len(media)}  {[Path(m).name for m in media]}")
    print(f"theme accent1/2: {accents}  fonts: "
          f"{major.group(1) if major else '?'} / {minor.group(1) if minor else '?'}")

    problems = []
    if totals["text"] < args.min_text:
        problems.append(f"only {totals['text']} text frames -- text may have been flattened")
    if args.expect_pictures is not None and totals["picture"] != args.expect_pictures:
        problems.append(f"expected {args.expect_pictures} pictures, found {totals['picture']}")
    if totals["picture"] and not media:
        problems.append("pictures reported but no media parts in the package")
    for p in problems:
        print("PROBLEM:", p)
    print("\nOK" if not problems else "\nFAILED")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
